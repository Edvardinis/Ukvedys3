import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

def generate_excel(tasks_df, faults_df, maintenance_df, expenses_df):
    output = io.BytesIO()
    wb = Workbook()
    
    ws_tasks = wb.active
    ws_tasks.title = "Darbų žurnalas"
    ws_faults = wb.create_sheet(title="Gedimų registras")
    ws_maint = wb.create_sheet(title="Profilaktinė priežiūra")
    ws_exp = wb.create_sheet(title="Išlaidos")
    
    def populate_sheet(ws, df):
        if df.empty:
            ws.append(["Nėra duomenų"])
            return
        headers = list(df.columns)
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for _, row in df.iterrows():
            ws.append([str(val) if val is not None else "" for val in row.values])
        for col in ws.columns:
            max_length = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_length + 3, 12)

    populate_sheet(ws_tasks, tasks_df)
    populate_sheet(ws_faults, faults_df)
    populate_sheet(ws_maint, maintenance_df)
    populate_sheet(ws_exp, expenses_df)
    
    wb.save(output)
    output.seek(0)
    return output.getvalue()

def generate_pdf_report(summary_data):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontSize=20, textColor=colors.HexColor('#2E7D32'), spaceAfter=15, alignment=1)
    h2_style = ParagraphStyle('H2Style', parent=styles['Heading2'], fontSize=14, textColor=colors.HexColor('#1A1A1A'), spaceBefore=10, spaceAfter=6)
    normal_style = ParagraphStyle('NormalStyle', parent=styles['Normal'], fontSize=10, textColor=colors.HexColor('#333333'), spaceAfter=4)
    
    story = []
    story.append(Paragraph("ŪKVEDŽIO VADOVO ATASKAITA", title_style))
    story.append(Paragraph(f"Ataskaitos data: {summary_data.get('date', '')}", normal_style))
    story.append(Spacer(1, 10))
    
    story.append(Paragraph("Pagrindiniai veiklos rodikliai", h2_style))
    data_table = [
        ["Rodiklis", "Reikšmė"],
        ["Gauti darbai", str(summary_data.get('total_tasks', 0))],
        ["Atlikti darbai", str(summary_data.get('completed_tasks', 0))],
        ["Atviri darbai", str(summary_data.get('open_tasks', 0))],
        ["Vėluojantys darbai", str(summary_data.get('overdue_tasks', 0))],
        ["Skubūs darbai (P1-P2)", str(summary_data.get('urgent_tasks', 0))],
        ["Registruoti gedimai", str(summary_data.get('total_faults', 0))],
        ["Neišspręsti gedimai", str(summary_data.get('open_faults', 0))],
        ["Sugaištos valandos iš viso", f"{summary_data.get('total_hours', 0):.1f} val."],
        ["Bendros išlaidos (€)", f"{summary_data.get('total_expenses', 0):.2f} €"]
    ]
    
    t = Table(data_table, colWidths=[250, 200])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#2E7D32')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,0), 6),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CCCCCC')),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F5F5F5')])
    ]))
    story.append(t)
    story.append(Spacer(1, 15))
    story.append(Paragraph("Pastabos ir rekomendacijos", h2_style))
    story.append(Paragraph(summary_data.get('notes_text', 'Nėra papildomų pastabų.'), normal_style))
    
    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()
